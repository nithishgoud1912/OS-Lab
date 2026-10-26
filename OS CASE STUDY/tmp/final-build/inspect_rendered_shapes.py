from pptx import Presentation
p=Presentation(r'C:\Users\HP\Documents\OS record\OS CASE STUDY\output\final\Case_Study_14_Race_Conditions_xv6_Final_v9.pptx')
for n in [2,5,8,9,11,13,14,15,16,17,18,20]:
 s=p.slides[n-1]; print('\nSLIDE',n)
 for sh in s.shapes:
  typ='TABLE' if sh.has_table else ('PIC' if sh.shape_type==13 else 'TEXT')
  txt=sh.text.replace('\n',' / ')[:110] if getattr(sh,'has_text_frame',False) and sh.text.strip() else ''
  sizes=[]
  if getattr(sh,'has_text_frame',False):
   for para in sh.text_frame.paragraphs:
    for run in para.runs:
     if run.font.size: sizes.append(round(run.font.size.pt,1))
  print(typ,tuple(round(v/914400,2) for v in [sh.left,sh.top,sh.width,sh.height]),sorted(set(sizes)),repr(txt))
