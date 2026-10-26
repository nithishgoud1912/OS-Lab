from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pathlib import Path
import csv, re, shutil, json
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(r'C:\Users\HP\Documents\OS record\OS CASE STUDY')
OUT=ROOT/'output'/'final'
BUILD=ROOT/'tmp'/'final-build'
DOCX_PATH=OUT/'Case_Study_14_Race_Conditions_xv6_Final_v5.docx'
FIG=ROOT/'member1_completed'/'figures'
FLOW_DIR=BUILD/'workflow-diagrams'
TERMINAL_DIR=BUILD/'terminal-stills'
RESULTS_CHART=BUILD/'results-comparison.png'

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.268)
sec.page_height=Inches(11.693)
sec.top_margin=Inches(.9)
sec.bottom_margin=Inches(.8)
sec.left_margin=Inches(1)
sec.right_margin=Inches(1)
sec.header_distance=Inches(.36)
sec.footer_distance=Inches(.45)
sec.different_first_page_header_footer=True

def set_font(run, size=12, bold=False, italic=False, color='000000', name='Times New Roman'):
    run.font.name=name
    run._element.rPr.rFonts.set(qn('w:eastAsia'),name)
    run.font.size=Pt(size)
    run.font.bold=bold
    run.font.italic=italic
    run.font.color.rgb=RGBColor.from_string(color)

normal=doc.styles['Normal']
normal.font.name='Times New Roman'
normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
normal.font.size=Pt(12)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.12
for style_name in ['Heading 1','Heading 2','Heading 3']:
    st=doc.styles[style_name]
    st.font.name='Times New Roman'
    st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    st.font.size=Pt(14)
    st.font.bold=True
    st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(10)
    st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.keep_with_next=True
if 'Code' not in [s.name for s in doc.styles]:
    code=doc.styles.add_style('Code',WD_STYLE_TYPE.PARAGRAPH)
else: code=doc.styles['Code']
code.font.name='Consolas'
code._element.rPr.rFonts.set(qn('w:eastAsia'),'Consolas')
code.font.size=Pt(10)
code.paragraph_format.left_indent=Inches(.18)
code.paragraph_format.right_indent=Inches(.1)
code.paragraph_format.space_before=Pt(2)
code.paragraph_format.space_after=Pt(2)
code.paragraph_format.line_spacing=1.0

def add_field(run, instr):
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    text=OxmlElement('w:instrText'); text.set(qn('xml:space'),'preserve'); text.text=instr
    separate=OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'),'separate')
    value=OxmlElement('w:t'); value.text='1'
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    run._r.extend([begin,text,separate,value,end])

def style_header_footer(header,footer):
    hp=header.paragraphs[0]
    hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.space_after=Pt(0)
    set_font(hp.add_run('OPERATING SYSTEMS CASE STUDY 14  |  XV6 SHARED DATA'),8.5,True,False,'596570')
    fp=footer.paragraphs[0]
    fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before=Pt(0)
    set_font(fp.add_run('Page '),8.5,color='596570')
    add_field(fp.runs[-1],'PAGE')
    set_font(fp.add_run(' of '),8.5,color='596570')
    add_field(fp.runs[-1],'NUMPAGES')
    set_font(fp.add_run('    |    Faculty signature: ____________________'),8.5,color='596570')
style_header_footer(sec.header,sec.footer)
sec.first_page_header.paragraphs[0].text=''
sec.first_page_footer.paragraphs[0].text=''

def p(text='',align=None,bold_prefix=None,style=None):
    para=doc.add_paragraph(style=style)
    if align is not None: para.alignment=align
    if bold_prefix and text.startswith(bold_prefix):
        set_font(para.add_run(bold_prefix),12,True)
        para.add_run(text[len(bold_prefix):])
    else: para.add_run(text)
    return para

def definition(label,text):
    para=doc.add_paragraph()
    para.paragraph_format.left_indent=Inches(.12)
    para.paragraph_format.right_indent=Inches(.08)
    para.paragraph_format.space_before=Pt(2)
    para.paragraph_format.space_after=Pt(7)
    para.paragraph_format.line_spacing=1.0
    ppr=para._p.get_or_add_pPr()
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'EDF3F8'); ppr.append(shd)
    borders=OxmlElement('w:pBdr')
    left=OxmlElement('w:left'); left.set(qn('w:val'),'single'); left.set(qn('w:sz'),'18'); left.set(qn('w:space'),'5'); left.set(qn('w:color'),'1F456E'); borders.append(left)
    ppr.append(borders)
    set_font(para.add_run(label),12,True,color='1F456E')
    set_font(para.add_run(': '+text),12)
    return para

def add_hyperlink(paragraph,text,url):
    rid=paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hyperlink=OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'),rid)
    run=OxmlElement('w:r'); rpr=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'0563C1'); rpr.append(color)
    underline=OxmlElement('w:u'); underline.set(qn('w:val'),'single'); rpr.append(underline)
    run.append(rpr)
    node=OxmlElement('w:t'); node.text=text; run.append(node)
    hyperlink.append(run); paragraph._p.append(hyperlink)

def heading(text,level=1):
    return doc.add_heading(text,level=level)

def set_cell_shading(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_cell_margins(cell,top=90,start=100,bottom=90,end=100):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr()
    tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def set_borders(table):
    tblPr=table._tbl.tblPr
    borders=tblPr.first_child_found_in('w:tblBorders')
    if borders is None: borders=OxmlElement('w:tblBorders'); tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=borders.find(qn('w:'+edge))
        if el is None: el=OxmlElement('w:'+edge); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),'D9D9D9')

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr()
    el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def prevent_row_split(row):
    trPr=row._tr.get_or_add_trPr()
    el=OxmlElement('w:cantSplit'); trPr.append(el)

def add_table(title,headers,rows,widths=None,font_size=9.5):
    cap=doc.add_paragraph()
    cap.paragraph_format.space_before=Pt(7); cap.paragraph_format.space_after=Pt(4); cap.paragraph_format.keep_with_next=True
    set_font(cap.add_run(title),10.5,True,color='000000')
    table=doc.add_table(rows=1,cols=len(headers))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    set_borders(table)
    all_rows=[headers]+rows
    for ri,vals in enumerate(all_rows):
        row=table.rows[0] if ri==0 else table.add_row()
        prevent_row_split(row)
        if ri==0: repeat_header(row)
        for ci,val in enumerate(vals):
            cell=row.cells[ci]
            if widths: cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            cell.text=''
            par=cell.paragraphs[0]
            par.paragraph_format.space_after=Pt(0)
            par.paragraph_format.line_spacing=1.0
            par.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci in (0,len(headers)-1) else WD_ALIGN_PARAGRAPH.CENTER
            run=par.add_run(str(val))
            set_font(run,font_size,bold=(ri==0),color=('FFFFFF' if ri==0 else '000000'))
            set_cell_shading(cell,'1F456E' if ri==0 else ('EDF3F8' if ri%2==0 else 'FFFFFF'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

def add_code(text):
    for line in text.strip('\n').splitlines():
        par=doc.add_paragraph(style='Code')
        par.paragraph_format.keep_together=True
        run=par.add_run(line if line else ' ')
        set_font(run,9,False,False,'202020','Consolas')

def add_figure(file,caption,width=6.15):
    par=doc.add_paragraph()
    par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before=Pt(5)
    par.paragraph_format.space_after=Pt(3)
    par.add_run().add_picture(str(file),width=Inches(width))
    cap=doc.add_paragraph()
    cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after=Pt(8)
    cap.paragraph_format.keep_with_next=False
    set_font(cap.add_run(caption),9.5,False,True,'333333')

def page_break():
    doc.add_page_break()

def make_flow_diagram(path,title,labels):
    width,height=1800,420
    image=Image.new('RGB',(width,height),'white')
    draw=ImageDraw.Draw(image)
    regular=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',30)
    bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',34)
    title_font=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',34)
    draw.text((56,32),title,font=title_font,fill='#12304A')
    margin=56; box_top=150; box_h=170; box_w=350
    gap=(width-2*margin-len(labels)*box_w)/(len(labels)-1)
    boxes=[]
    for i,label in enumerate(labels):
        x=round(margin+i*(box_w+gap))
        draw.rounded_rectangle((x,box_top,x+box_w,box_top+box_h),radius=6,fill='#F1F3F5',outline='#1F456E',width=4)
        lines=label.split('\n')
        heights=[draw.textbbox((0,0),line,font=bold)[3]-draw.textbbox((0,0),line,font=bold)[1] for line in lines]
        total=sum(heights)+14*(len(lines)-1)
        y=box_top+(box_h-total)/2
        for line,line_h in zip(lines,heights):
            bb=draw.textbbox((0,0),line,font=bold); text_w=bb[2]-bb[0]
            draw.text((x+(box_w-text_w)/2,y),line,font=bold,fill='#17202A')
            y+=line_h+14
        boxes.append((x,x+box_w))
    mid=box_top+box_h//2
    for (left,right),(next_left,next_right) in zip(boxes,boxes[1:]):
        y=mid
        start=right+7; end=next_left-11
        draw.line((start,y,end,y),fill='#607487',width=5)
        draw.polygon([(end,y),(end-18,y-12),(end-18,y+12)],fill='#607487')
    image.save(path)

def make_terminal_still(cast_name,out_name,phase_title):
    TERMINAL_DIR.mkdir(parents=True,exist_ok=True)
    source=ROOT/'member1_completed'/'evidence'/cast_name
    events=[json.loads(line) for line in source.read_text(encoding='utf-8').splitlines()]
    raw=''.join(event[2] for event in events[1:] if len(event)>2 and event[1]=='o')
    raw=re.sub(r'\x1b\][^\x07]*(?:\x07|\x1b\\)','',raw)
    raw=re.sub(r'\x1b\[[0-?]*[ -/]*[@-~]','',raw)
    lines=[line.rstrip() for line in raw.replace('\r','').splitlines()]
    while lines and not lines[0]: lines.pop(0)
    while lines and not lines[-1]: lines.pop()
    W,H=1800,max(650,265+42*len(lines))
    image=Image.new('RGB',(W,H),'#FFFFFF'); draw=ImageDraw.Draw(image)
    font=ImageFont.truetype('C:/Windows/Fonts/consola.ttf',28)
    label=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',27)
    small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',20)
    draw.text((56,34),phase_title,font=label,fill='#17365D')
    draw.text((56,76),f'Static render from saved terminal recording: {cast_name}',font=small,fill='#53616E')
    draw.rounded_rectangle((48,120,W-48,H-42),radius=10,fill='#F2F4F6',outline='#AAB4BF',width=2)
    draw.rounded_rectangle((48,120,W-48,166),radius=10,fill='#DDE3E8',outline='#AAB4BF',width=2)
    for i,color in enumerate(['#7D8790','#A4ADB5','#C1C8CE']):
        x=75+i*27; draw.ellipse((x,136,x+13,149),fill=color)
    y=187
    max_lines=(H-235)//42
    for line in lines[:max_lines]:
        draw.text((76,y),line,font=font,fill='#17202A')
        y+=42
    image.save(TERMINAL_DIR/out_name)
    return TERMINAL_DIR/out_name

def make_results_chart(path):
    W,H=1800,900
    image=Image.new('RGB',(W,H),'#FFFFFF'); draw=ImageDraw.Draw(image)
    title=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',54)
    heading_font=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',44)
    body=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',42)
    value_font=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',38)
    ink='#17202A'; muted='#53616E'; border='#AAB4BF'
    colors=[('#1F456E','Corruption'),('#697786','Panic / trap'),('#9AA5AF','Test failure'),('#D8DEE3','Timeout')]
    draw.text((68,38),'Unlocked phase outcomes (10 attempts per dataset)',font=title,fill='#17365D')
    data=[('Original set',[3,4,2,1]),('C-runner validation',[1,6,2,1])]
    left=400; bar_w=1230; bar_h=90; gap=135; y0=200
    for idx,(label,counts) in enumerate(data):
        y=y0+idx*gap
        draw.text((68,y+18),label,font=heading_font,fill=ink)
        x=left
        for count,(color,_) in zip(counts,colors):
            width=bar_w*count/10
            draw.rectangle((x,y,x+width,y+bar_h),fill=color,outline='#FFFFFF',width=3)
            text=str(count); bb=draw.textbbox((0,0),text,font=value_font); tw=bb[2]-bb[0]
            draw.text((x+(width-tw)/2,y+25),text,font=value_font,fill=('#FFFFFF' if color in ('#1F456E','#697786') else ink))
            x+=width
    legend_y=510; x=68
    for color,label in colors:
        draw.rectangle((x,legend_y,x+30,legend_y+30),fill=color)
        draw.text((x+42,legend_y-2),label,font=body,fill=ink)
        x+=430
    draw.line((68,600,W-68,600),fill=border,width=2)
    draw.text((68,640),'Locked: 10/10 PASS     Restored: 10/10 PASS',font=heading_font,fill='#17365D')
    draw.text((68,705),'Counts are primary classifications from two separate 30-attempt datasets.',font=body,fill=muted)
    draw.text((68,750),'A timeout is not a deadlock diagnosis; a pass is not proof of race freedom.',font=body,fill=muted)
    image.save(path)

TERMINAL_A01=make_terminal_still('A01-locked.cast','A01-locked.png','LOCKED BASELINE • A01')
TERMINAL_B02=make_terminal_still('B02-unlocked.cast','B02-unlocked.png','UNLOCKED RUN • B02')
TERMINAL_C01=make_terminal_still('C01-restored.cast','C01-restored.png','RESTORED LOCKS • C01')
make_results_chart(RESULTS_CHART)

FLOW_DIR.mkdir(parents=True,exist_ok=True)
ALLOC_FLOW=FLOW_DIR/'allocation-path.png'
PHASE_FLOW=FLOW_DIR/'runner-phases.png'
make_flow_diagram(ALLOC_FLOW,'Guest allocation path',[
    'race_test.c\nsbrk(bytes)','sys_sbrk()\ngrowproc()','uvmalloc()\nkalloc()','kmem.freelist\nkmem.lock'])
make_flow_diagram(PHASE_FLOW,'Host experiment and recovery',[
    'A: LOCKED\noriginal calls','B: UNLOCKED\n4 calls disabled','C: RESTORED\nsaved source rebuilt','VERIFY\nhash + diff + CSV'])

# Cover page.
for _ in range(2): p('',WD_ALIGN_PARAGRAPH.CENTER)
for text,size in [('CHAITANYA BHARATHI INSTITUTE OF TECHNOLOGY',16),('Department of Computer Science and Engineering',14),('Operating Systems',13)]:
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(8)
    set_font(para.add_run(text),size,True if size>=14 else False)
p('',WD_ALIGN_PARAGRAPH.CENTER)
para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.paragraph_format.space_before=Pt(24)
set_font(para.add_run('CASE STUDY 14'),18,True,color='17365D')
para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
para.paragraph_format.space_after=Pt(18)
set_font(para.add_run('Race Condition in XV6 Shared Data'),22,True,color='17365D')
para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_font(para.add_run('Study of shared physical-memory free-list synchronization'),13,False,color='465A6E')
p('',WD_ALIGN_PARAGRAPH.CENTER)
add_table('Team members', ['Name','Roll number'], [
    ['L Nithish Kumar Goud','160124733327'],
    ['G Ishan','160124733322'],
    ['Santhosh','160124733332'],
    ['Venkat','160124733340']], [3.6,2.0],10.5)
for text in ['Submitted to: ______________________________________________','Faculty / designation: _______________________________________','Semester / academic year: ___________________________________','Submission date: ____________________________________________']:
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.paragraph_format.space_after=Pt(8)
    set_font(para.add_run(text),11)

page_break()
heading('Certificate and team responsibility',1)
p('Team name: Case Study 14 — XV6 Shared Data')
p('Problem statement: Study race conditions when multiple processes or CPUs access shared kernel data without synchronization, and analyze how xv6 locks prevent inconsistent updates.')
add_table('Table 1. Supplied four-member responsibility allocation', ['Member','Lead responsibility in division plan','Share'], [
    ['L Nithish Kumar Goud (160124733327)','Technical implementation, experiment, evidence and results', '25%'],
    ['G Ishan (160124733322)','Lock-removal analysis and race timelines','25%'],
    ['Santhosh (160124733332)','Spinlock and critical-section explanation','25%'],
    ['Venkat (160124733340)','Foundations, shared-data inventory and report integration','25%']], [2.05,3.85,.55],9.5)
p('Faculty name: ______________________________  Signature: ______________________________  Date: ______________')

page_break()
heading('Index',1)
index_rows=[
('Case statement, objectives, activities and expected outcome','4'),
('1. What is a race condition?','4'),
('2. Problems generated by race conditions','5'),
('Existing system and background','5'),
('3. Proposed solution','6'),
('4. What is synchronization?','6'),
('5. Types of synchronization','6'),
('6. How xv6 protects shared data','7'),
('Working flow','8'),
('7. C files, WSL commands and experiment sequence','8'),
('7.6. Kernel allocator critical sections','12'),
('7.7. Spinlock ownership and memory ordering','13'),
('7.8. Guest workload gate and checker','14'),
('7.9. Host runner phases and restoration','15'),
('8. What the locks prevent, and possible unlocked outcomes','15'),
('9. Comparison of results','16'),
('10. Conclusion, applications and future scope','19'),
('References','20')]
index_table=add_table('Table 2. Report contents', ['Section','Page'], index_rows,[5.3,.65],9.5)
p('')

page_break()
heading('Case Study 14: Problem Statement, Objectives, Activities and Expected Outcome',1)
p('Problem Statement: Study how race conditions can occur when multiple processes or CPUs access shared kernel data without proper synchronization. Analyse how XV6 prevents inconsistent updates using locking mechanisms.')
p('The case uses the physical-memory free list as the shared kernel data structure. The study follows the access paths and the lock associated with the structure, reasons about possible outcomes when the lock calls are removed, and compares locked, unlocked, and restored behavior using the same user workload.')
heading('Objectives',2)
p('The objectives are to understand race conditions and critical sections; identify shared kernel data structures in xv6; study the need for mutual exclusion; observe the effect of concurrent updates; and explain how locks prevent inconsistent results.')
heading('Activities',2)
p('The completed source review identifies the allocator free list and its lock. The experiment used four CPUs and a coordinated multi-process workload. It compared original locking, a temporary removal of four allocator lock calls, and restoration. The report distinguishes possible source-level interleavings from events directly proven by the logs.')
heading('Expected outcome',2)
p('Students should be able to explain how race conditions arise and how synchronization mechanisms provide safe access to shared kernel data.')

heading('1. What is a race condition?',1)
definition('Race condition','concurrent operations access shared data, and the result depends on how their steps interleave.')
p('A critical section reads or changes shared state; mutual exclusion lets only one participant execute that section at a time.')
p('The motivation is to connect synchronization theory to a real kernel resource. xv6 manages memory on multiple CPUs, so a list update that appears simple in source can become unsafe when operations overlap.')
p('The xv6 race is a multi-step pointer update. kalloc() reads the free-list head P, obtains P->next, advances kmem.freelist, and returns P. If CPU 1 reads the same P before CPU 0 advances the head, both calls can return the same physical page. Each process can then write through a different virtual address that maps to that one page.')
add_table('Table 3. Possible duplicate allocation without kmem.lock', ['Step','CPU 0: kalloc()','CPU 1: kalloc()','Shared head'],[
    ['1','r0 = kmem.freelist = P','', 'P'],
    ['2','','r1 = kmem.freelist = P','P'],
    ['3','kmem.freelist = P->next','', 'P->next'],
    ['4','','kmem.freelist = P->next','P->next'],
    ['5','return P','return P','Both callers own P']], [.55,1.65,1.65,1.45],9.1)
p('A second dangerous path occurs in kfree(): two CPUs can each link a returned page to the same stale head, then overwrite one another’s head update. The resulting schedule may lose a page or break the list. These are source-level possibilities; the logs do not reconstruct which instructions interleaved in a given attempt.')

heading('2. Problems generated by race conditions',1)
p('When operations that preserve a shared-data invariant interleave without the required lock, one update can overwrite another or a reader can observe a partly updated structure. In the page allocator, possible consequences include returning the same page twice, losing a free page, corrupting a next pointer, failing an allocation, or trapping in the kernel.')
add_table('Table 4. Possible problems and their effect', ['Problem','Possible allocator cause','Possible effect'],[
    ['Lost update','One CPU overwrites a newer free-list head','A page becomes unreachable from the list'],
    ['Duplicate allocation','Two CPUs read the same head before either removes it','Two callers may use the same physical page'],
    ['Pointer or list corruption','Conflicting writes damage a next pointer or head','Later allocator access may fail or trap'],
    ['Allocation/test failure','Allocator returns bad state or cannot satisfy request','A process reports failure or kernel activity stops'],
    ['Deadlock confusion','A run does not finish within a timeout','Timeout alone does not show why it stopped']], [1.25,2.5,2.1],9)
p('A race is a correctness problem. A deadlock is a progress problem in which participants wait indefinitely for conditions that cannot be met. The experiment records a host timeout as TIMEOUT, not as a diagnosed deadlock. A panic or trap also does not prove which exact instruction interleaving occurred.')

heading('Existing System and Background',1)
p('Standard xv6 uses one global kmem.lock to protect its physical-page free list. This coarse-grained design is simple and preserves list consistency, but every CPU requesting or freeing a page shares the same lock and may spin while another CPU owns it. The MIT xv6 locking chapter discusses this contention and per-CPU free lists as a possible design that reduces sharing, at the cost of refill or page-stealing logic.')
p('A lock-free compare-and-swap (CAS) loop is another design option. It avoids waiting on one lock but needs careful handling of pointer lifetime and ABA, where a location changes from A to B and back to A so a CAS cannot detect the intervening change. Per-CPU lists and lock-free allocation were background alternatives only; neither was implemented or benchmarked here.')
p('Other shared xv6 data reviewed at source level: process state (p->lock), PID allocation (pid_lock), parent/child wait state (wait_lock), buffer-cache metadata (bcache.lock), and timer ticks (tickslock). These structures were not experimentally modified.')

heading('3. Proposed Solution: protect the complete update',1)
p('A shared invariant needs one consistent synchronization rule. For the allocator, every operation that reads or changes the free-list head and links must use kmem.lock. The lock must cover dependent reads as well as writes; protecting only the final assignment would leave the same stale-read race.')
p('A spinlock provides mutual exclusion by giving one CPU ownership of the critical section while other CPUs wait. Atomic exchange provides the cross-CPU ownership test. Acquire ordering ensures following protected operations occur after acquisition, while release ordering publishes the completed updates before another CPU acquires the lock. xv6 also disables local interrupts while holding a spinlock to avoid same-CPU interrupt re-entry. Interrupt disabling by itself cannot exclude a second CPU.')
p('The tested source already contains this solution. The experiment removes the acquire and release calls temporarily, keeps the workload and list operations unchanged, then restores the original allocator byte for byte and rebuilds it. The experiment therefore demonstrates an existing xv6 mechanism rather than introducing a new one.')

heading('4. What is synchronization?',1)
definition('Synchronization','coordination of concurrent operations so shared state remains correct.')
p('Mutual exclusion is one form. Other mechanisms coordinate ordering, wait for a condition, count available resources, or mark a common phase boundary.')
add_table('Table 5. Spinlock operation in the checked xv6 revision', ['Stage','Operation','Purpose'],[
    ['Before acquire','push_off()','Disable local interrupts and track nesting'],
    ['Acquire','Atomic exchange with acquire ordering','Claim exclusive ownership across CPUs'],
    ['Critical section','Update kmem.freelist and page links','Preserve the allocator invariant'],
    ['Release','Atomic store with release ordering','Publish updates and make the lock available'],
    ['After release','pop_off()','Restore prior interrupt state when nesting reaches zero']], [1.25,2.45,2.15],9)
p('In kernel/spinlock.c, acquire uses __atomic_exchange_n with __ATOMIC_ACQUIRE and release uses __atomic_store_n with __ATOMIC_RELEASE. push_off and pop_off manage local interrupt state and nesting. The atomic exchange supplies cross-CPU mutual exclusion. Local interrupt disabling alone does not provide it.')

heading('5. Types of synchronization',1)
add_table('Table 6. Common synchronization mechanisms', ['Mechanism','How it coordinates work','Typical use or limitation'],[
    ['Mutex','One thread owns a protected region at a time','Blocking mutual exclusion in thread libraries'],
    ['Spinlock','One CPU owns the region; others wait actively','Short kernel critical sections; consumes CPU while waiting'],
    ['Semaphore','A counter tracks available permits','Bounded resources or event counts'],
    ['Condition variable / monitor','A thread sleeps until a predicate can be checked again under a mutex','Wait for a state change without busy waiting'],
    ['Reader-writer lock','Many readers or one writer','Read-heavy data; writer fairness and complexity matter'],
    ['Barrier','Participants wait at a phase boundary','Parallel computation phases'],
    ['Atomic operation','One indivisible machine update','Single values; not automatically a whole-object invariant']], [1.45,2.55,1.85],8.8)
p('The project measured only xv6’s physical allocator spinlock. The other mechanisms are background concepts and were not implemented or benchmarked in this study.')

heading('6. How xv6 protects the shared data',1)
p('The checked source revision is f780d7850765d74cbdcd1fff4b822d8829f7c9bd. In kernel/kalloc.c, kmem.freelist points to unused physical pages and kmem.lock protects updates to the list. kalloc removes the first page while holding the lock. kfree links a released page at the head while holding the same lock.')
add_table('Table 7. xv6 files examined and their role', ['File','Role in the path','Status in this study'],[
    ['kernel/kalloc.c','Defines kmem.freelist and kmem.lock; implements kalloc and kfree','Lock calls removed and restored for the measured experiment'],
    ['kernel/spinlock.c / spinlock.h','Implements initlock, acquire, release, holding, push_off, pop_off','Read to explain the existing lock'],
    ['user/race_test.c','Runs the four-child memory verification workload inside xv6','Included in source archive and unchanged across phases'],
    ['kernel/sysproc.c','sys_sbrk dispatches memory growth or shrink','Source-level path review'],
    ['kernel/proc.c','growproc manages process address-space size','Source-level path review'],
    ['kernel/vm.c','uvmalloc/uvmdealloc and uvmunmap allocate or release pages','Source-level path review'],
    ['user/ulib.c','sbrk wrapper invokes the system call','Source-level path review'],
    ['member1_tools.c','Host-side C runner builds, runs, records, restores, packages','C implementation used for fresh validation']],
    [1.55,2.85,1.45],8.6)
p('The path for growth is user/ulib.c sbrk, kernel/sysproc.c sys_sbrk, kernel/proc.c growproc, kernel/vm.c uvmalloc, then kernel/kalloc.c kalloc. Shrink follows growproc to uvmdealloc and uvmunmap, then kfree. Fork and exit also exercise the allocator, but the stress program’s direct path is through sbrk.')

heading('Working flow',1)
add_figure(ALLOC_FLOW,'Figure 1. Guest allocation and release path from sbrk() to the shared allocator.',6.15)
add_figure(PHASE_FLOW,'Figure 2. Locked, unlocked, and restored experiment phases with final checks.',6.15)
p('Each attempt boots a fresh QEMU guest. The host runner captures console output, classifies one primary result, and appends the run ID and evidence path to results.csv.')

heading('7. C files, WSL commands and experiment sequence',1)
p('The case study has two project C programs with different execution environments. user/race_test.c is a guest user program compiled by the xv6 RISC-V Makefile. member1_tools.c is a host program compiled by Ubuntu cc under WSL. kernel/kalloc.c is xv6 kernel C and contains the existing allocator. The lock-removal variant changes only four acquire/release calls in that source file.')
heading('7.1 user/race_test.c: guest workload',2)
p('The parent creates a pipe, forks four children, and waits until all children exist before releasing them. Each child repeats 200 rounds. In one round it requests 32768 bytes through sbrk, fills the allocation with its distinct byte tag, waits briefly in user space, and checks every byte. It then shrinks the allocation using sbrk(-32768). The parent waits for all children and checks setup and exit statuses before printing PASS or FAIL.')
p('The pipe gate encourages concurrent activity. Each request covers eight 4096-byte pages, so all four children repeatedly exercise the kernel allocator. The test checks the contents visible in each user buffer. It may miss a page lost from the free list or corruption outside those buffers. volatile forces verification reads in this program; it is not a lock and does not synchronize kernel access.')
heading('7.2 member1_tools.c: host runner',2)
p('The C host helper automates builds, QEMU runs, timed console capture, experiment phases, restoration, and result packaging. It preserves console logs and terminal recordings, writes per-attempt CSV rows, checks the restored allocator, and can replay a saved transcript. The following table maps every named helper in the source to its role.')
add_table('Table 8. member1_tools.c function map', ['Function','Responsibility'],[
    ['signal_stop','Records a stop request so the runner can shut down and restore during handled interruption'],
    ['fail','Reports a fatal error and terminates the current operation'],
    ['path_join','Builds a checked path from directory and filename parts'],
    ['load / save','Reads and writes complete files used for snapshots, logs and generated summaries'],
    ['execute','Starts a child process, captures output, enforces time limits, and reports status'],
    ['checked','Runs a required command and fails if it returns an unsuccessful status'],
    ['restore','Copies the saved allocator back and rebuilds the restored kernel'],
    ['json_string','Escapes text for JSON output'],
    ['csv_text','Quotes and escapes CSV text fields'],
    ['message_line','Appends one output line to a captured transcript'],
    ['build','Runs and records the xv6 build commands'],
    ['record_command','Saves the command and its output for the evidence archive'],
    ['run','Coordinates preflight, locked attempts, lock removal, unlocked attempts, restoration, and restored attempts'],
    ['csv_row','Writes one attempt record with phase, status, elapsed time, and evidence filenames'],
    ['compare_names','Sorts file names for stable packaging and manifest generation'],
    ['manifest_tree','Walks the evidence tree and creates file/checksum records'],
    ['package','Checks the results, writes summaries/checksums, and creates the handoff archive'],
    ['main','Parses the run, package, and replay command-line modes']], [1.45,4.75],8.5)
p('A handled Ctrl+C requests the runner shutdown path, which restores and rebuilds the allocator. A power loss or forced termination such as SIGKILL cannot run that cleanup. The saved final restoration check and empty allocator diff are the evidence that the checked-in working source was restored.')

heading('7.3 Environment and fixed workload',2)
add_table('Table 9. Recorded experiment configuration', ['Item','Recorded value'],[
    ['Source','xv6-riscv commit f780d7850765d74cbdcd1fff4b822d8829f7c9bd'],
    ['Host and guest','Ubuntu 24.04 under WSL2; RISC-V xv6 guest'],
    ['Emulator and compiler','QEMU 8.2.2; RISC-V GCC 13.3.0'],
    ['Guest resources','4 CPUs; 128 MiB RAM'],
    ['Workload','4 children; 8 pages per child allocation; 200 rounds'],
    ['Attempts per phase','10 locked; 10 unlocked; 10 restored'],
    ['Run isolation','Fresh QEMU guest for each attempt'],
    ['Limits','30 seconds to reach shell; 60 seconds after stress command'],
    ['Instrumentation','No artificial kernel delay; no kernel timing instrumentation']], [1.65,4.55],9)
add_figure(TERMINAL_B02,'Figure 3. Static terminal still reconstructed from saved B02 recording; output includes corruption and kerneltrap.',5.9)

heading('7.4 WSL setup and command sequence',2)
p('On this PC, Ubuntu 24.04, QEMU, and the RISC-V toolchain were already available. The following PowerShell command starts the distribution. The package-install commands are included for a new Ubuntu installation and were not part of the measured attempts.')
add_code(r'''# PowerShell
wsl -d Ubuntu-24.04''')
add_code(r'''# New Ubuntu setup only
sudo apt update
sudo apt install git build-essential qemu-system-misc gcc-riscv64-linux-gnu binutils-riscv64-linux-gnu bc''')
p('For a new checkout, use an unused directory. Do not clone over the completed ~/xv6-member1 checkout. The revision below is the fixed commit used by the case study.')
add_code(r'''git clone https://github.com/mit-pdos/xv6-riscv.git ~/xv6-member1
cd ~/xv6-member1
git checkout --detach f780d7850765d74cbdcd1fff4b822d8829f7c9bd
git switch -c member1-race-study''')
p('The completed checkout includes user/race_test.c and a UPROGS entry in Makefile. The entry uses a literal tab before the target.')
add_code("# Add to the existing UPROGS list in Makefile\n\t$U/_race_test\\")
p('The ordinary build and sanity sequence boots xv6, checks the known revision, runs the saved OS sanity command, and then invokes the stress program from the xv6 shell.')
add_code(r'''cd ~/xv6-member1
git rev-parse HEAD
make -j4 kernel/kernel fs.img
make qemu CPUS=4
# At the xv6 shell:
echo hello xv6
usertests -q
race_test
# Exit QEMU with Ctrl+A, then X''')
p('The growth path begins at sbrk and ends at kalloc. Shrinking follows uvmdealloc and uvmunmap to kfree. The test is run as a user program and does not call kernel functions directly.')

heading('7.5 C runner commands and the three measured phases',2)
p('The host runner is compiled with the Ubuntu C compiler. The following run command is a reproducible fresh-run example that writes to a new evidence directory. It is not presented as the exact historical output path. The runner refuses a directory that already exists and expects the original allocator to match Git HEAD.')
add_code(r'''cd '/mnt/c/Users/HP/Documents/OS record/OS CASE STUDY'
cc -std=c11 -Wall -Wextra -Werror -O2 member1_tools.c -o /tmp/member1_tools
/tmp/member1_tools run /home/nithish/xv6-member1 /home/nithish/member1-new-evidence full''')
p('Run mode performs sanity and locked smoke checks before the measured attempts. It then executes the same workload across the three phases. A fresh guest boots for every attempt.')
add_table('Table 10. Phase sequence handled by the C runner', ['Phase','Allocator condition','Recorded run IDs','Restoration action'],[
    ['A: locked','Original kmem.lock calls present','A01–A10','Save source and establish baseline'],
    ['B: unlocked','Comment out acquire/release in kalloc and kfree','B01–B10','Keep allocator data operations and workload unchanged'],
    ['C: restored','Original allocator restored byte for byte','C01–C10','Verify checksum, empty final diff, rebuild and repeat']], [1.0,2.0,1.05,2.2],8.8)
p('The runner has separate package and replay modes. Replay prints a saved transcript, while package validates the archived CSV, builds summaries and checksums, and creates the compressed handoff.')

section76=heading('7.6 Kernel allocator critical sections',2)
section76.paragraph_format.page_break_before=True
p('The experiment targets the linked list in kernel/kalloc.c. kmem.freelist points to the first free physical page, and each free page stores its next pointer in struct run. A page can be removed or inserted only while the list head and the dependent next pointer are protected by the same lock.')
add_code(r'''/* kalloc(): detach one page from the list */
acquire(&kmem.lock);
r = kmem.freelist;
if (r)
  kmem.freelist = r->next;
release(&kmem.lock);

/* kfree(): prepend one page to the list */
acquire(&kmem.lock);
r->next = kmem.freelist;
kmem.freelist = r;
release(&kmem.lock);''')
p('The lock protects the transfer between the shared list and a caller. kalloc() fills the removed page after releasing the lock because the page is no longer reachable by another allocator. kfree() fills the page before linking it, so no other allocator can obtain that page during the fill. The critical section protects list membership, not later writes made by the process that owns an allocated page.')

page_break()
heading('7.7 Spinlock ownership and memory ordering',2)
p('The checked kernel/spinlock.c implementation combines local interrupt control with an atomic cross-CPU ownership operation. The atomic exchange is the point where one CPU claims the lock. Acquire ordering keeps protected memory operations after that claim; release ordering publishes completed writes before another CPU obtains the lock.')
add_code(r'''push_off();
if (holding(lk))
  panic("acquire");
while (__atomic_exchange_n(&lk->locked, 1, __ATOMIC_ACQUIRE) != 0)
  ;
lk->cpu = mycpu();
...
__atomic_store_n(&lk->locked, 0, __ATOMIC_RELEASE);
pop_off();''')
add_table('Table 11. What each spinlock operation protects', ['Operation','Effect in this xv6 revision','Boundary'],[
    ['push_off()','Disables local interrupts and tracks nesting','Does not exclude another CPU'],
    ['Atomic exchange','Changes unlocked to owned as one operation','Other CPUs spin while owned'],
    ['Acquire ordering','Keeps critical-section accesses after ownership','Does not choose the data invariant'],
    ['Release store','Makes protected updates visible before unlock','Does not guarantee fairness'],
    ['pop_off()','Restores interrupt state when nesting returns to zero','Runs after ownership is released']], [1.25,2.7,1.7],8.8)
p('holding() checks whether the current CPU owns the lock, while cpu and locked fields support ownership checks and diagnostics. The lock is intentionally short: the allocator releases it before page-filling work in kalloc(). Removing acquire/release also removes the interrupt-handling effects of those calls when no outer spinlock is held, so the measured change is not an isolated test of atomic ownership alone.')

page_break()
heading('7.8 Guest workload gate and memory checker',2)
p('user/race_test.c creates four children and holds each child at a pipe read until all forks succeed. The parent then writes four tokens. This narrows the start-time spread without injecting a delay into kernel code. Each worker uses an eight-page allocation, writes a distinct byte tag, waits briefly in user space, reads every byte again, and shrinks the address space.')
add_code(r'''int pid = fork();
if (pid == 0) {
  int received = read(gate[0], &token, 1);
  if (received != 1)
    exit(2);
  exit(worker(id));
}
...
write(gate[1], tokens, NCHILD);

char *allocation = sbrk(bytes);
volatile uchar *memory = (volatile uchar *)allocation;
for (int offset = 0; offset < bytes; offset++)
  memory[offset] = tag;
/* second pass reports a mismatch before sbrk(-bytes) */''')
p('The explicit user-page workload contributes a lower bound of 4 children × 200 rounds × 8 pages = 6,400 page allocations plus 6,400 releases, or 12,800 kalloc()/kfree() calls if every round completes. This is not a measured count or a total: page-table pages, fork, pipes, and other kernel work add calls. The test checks whether processes retain distinct data in their user buffers; passing does not detect every lost or duplicate page.')
p('The parent treats pipe or fork setup errors and every non-zero wait() status as failure. It reports PASS only when all four children start and exit successfully. volatile forces the verification loop to read memory again; it is not an inter-process barrier and does not make the kernel allocator safe.')

page_break()
heading('7.9 Host runner phases and restoration',2)
p('member1_tools.c is a host C program compiled inside Ubuntu under WSL. Its run mode saves the original allocator, verifies the fixed revision and workload files, builds xv6, launches QEMU through pipes, polls console output, applies time limits, and records a primary outcome with the raw transcript. The runner starts one fresh guest per attempt and records the two evidence datasets in separate directories.')
add_code(r'''must_restore = 1;
atexit(restore);
...
if (edits != 4)
  exit(1);
save(allocator, variant, dest);
...
else if (phase == 2)
  save(allocator, original, original_size);
...
restore();
sha256sum -c allocator-original.sha256
git diff --exit-code HEAD -- kernel/kalloc.c''')
p('Before the unlocked phase, the runner comments exactly four acquire/release calls in kalloc.c and saves the resulting diff. Before the restored phase, it writes back the saved allocator bytes and rebuilds kernel/kernel and fs.img. It verifies the saved source checksum and requires an empty final diff. atexit(restore) covers handled process termination; it cannot recover the file after power loss or forced termination. A final verification therefore remains part of the recorded workflow.')
p('For each QEMU run, the host captures console text to a .txt transcript and a timed terminal recording, then appends the attempt ID, phase, status, elapsed time, and evidence filename to results.csv. Elapsed seconds begin when the host recognizes the xv6 shell prompt immediately before sending race_test; they include result detection and a short output-collection window, not boot time or pure guest execution time.')
p('The C-runner selects one primary result from accumulated output in this order: CORRUPTION, PANIC_OR_TRAP, TEST_FAILURE, then PASS. If the host time limit expires, TIMEOUT overrides the detected text. Thus B02 remains CORRUPTION despite a later kerneltrap, while B10 is TIMEOUT despite earlier trap diagnostics. Neither classification establishes the precise instruction interleaving.')
p('From the Windows-mounted project directory, the helper can replay a saved attempt and package the evidence archive:')
add_code(r'''cd '/mnt/c/Users/HP/Documents/OS record/OS CASE STUDY'
/tmp/member1_tools replay member1_completed/evidence/B02-unlocked.txt
/tmp/member1_tools package member1_completed''')
p('The original results remain in member1_completed/evidence/ and the independent C-runner set in member1_completed/c_validation/. The two 30-attempt datasets are not merged.')

page_break()

heading('8. What the locks prevent, and possible unlocked outcomes',1)
p('A lock does not prevent concurrency; it makes the protected operations run one at a time with respect to the same lock. That ordering preserves the free-list invariant. Without the lock, a harmful schedule is possible even if the two individual pointer assignments are each machine-word operations.')
add_table('Table 12. Possible kalloc interleaving without kmem.lock', ['Step','CPU 0','CPU 1','Possible state'],[
    ['1','Read freelist head P','', 'Head is P'],
    ['2','','Read freelist head P','Both hold stale P'],
    ['3','Set head to P.next','', 'Head advances'],
    ['4','','Set head to P.next','Head remains advanced'],
    ['5','Return P','Return P','Both callers may receive P']], [.5,1.65,1.65,2.3],8.8)
add_table('Table 13. Possible kfree lost-page schedule without kmem.lock', ['Step','CPU 0 frees A','CPU 1 frees B','Possible state'],[
    ['1','Read old head F','', 'Head is F'],
    ['2','','Read old head F','Both read F'],
    ['3','Set A.next to F','', 'A points to F'],
    ['4','','Set B.next to F','B points to F'],
    ['5','Set head to A','', 'Head becomes A'],
    ['6','','Set head to B','A may be lost from reachable list']], [.5,1.65,1.65,2.3],8.8)
p('These schedules are source-level explanations, not captured instruction traces. The recorded corruption, test failures, panics, and traps demonstrate that removing allocator locking produced unsafe outcomes under the tested workload. They do not identify which exact schedule caused each run.')
heading('9. Comparison of results',1)
p('Two datasets are preserved. The original run set was recorded before the host helper was converted to C and is stored under member1_completed/evidence/. The fresh validation run set was recorded by member1_tools.c and is stored under member1_completed/c_validation/. Each dataset has 30 attempts, with 10 attempts in each phase. Counts below are not merged.')
add_table('Table 14. Outcomes by independent dataset', ['Dataset','Locked phase','Unlocked phase','Restored phase'],[
    ['Original observations','10/10 PASS','3 corruption; 4 panic/trap; 2 test failure; 1 timeout','10/10 PASS'],
    ['C-runner validation','10/10 PASS','1 corruption; 6 panic/trap; 2 test failure; 1 timeout','10/10 PASS']], [1.35,1.05,2.5,1.05],8.5)
p('In both datasets, every attempt reached the xv6 shell and started race_test. The locked and restored phases passed in all ten attempts. The unlocked phase produced ten non-PASS outcomes in each set. In the original set, B10 printed trap diagnostics before the host timeout and is counted once as TIMEOUT. Some multicore output is interleaved.')
p('The observations support the importance of allocator synchronization for this revision and workload. A PASS is one observed successful invocation, not proof that no harmful schedule exists. A timeout does not prove deadlock. The test may miss lost pages or corruption outside its user buffers. The elapsed times include output collection and are not performance benchmarks.')
add_figure(RESULTS_CHART,'Figure 4. Comparison of primary outcomes across the two separate datasets.',5.9)
heading('9.1 Original observations: all 30 attempts',2)
p('The following table reproduces every primary outcome and log filename from evidence/results.csv. The original unlocked set contains three corruption records, four panic/trap records, two test failures, and one timeout.')
def result_rows(file,prefix):
    with open(file,newline='',encoding='utf-8') as f:
        data=list(csv.DictReader(f))
    return [[r['attempt_id'],r['phase'],r['outcome'],r['elapsed_seconds'],prefix+'/'+r['transcript']] for r in data]
orig_rows=result_rows(ROOT/'member1_completed'/'evidence'/'results.csv','evidence')
add_table('Table 15. Original dataset, attempt-level results', ['Run ID','Phase','Primary outcome','Elapsed (s)','Raw log'],orig_rows,[.62,.82,1.35,.78,2.6],9.0)
heading('9.2 C-runner validation: all 30 attempts',2)
p('This independent table reproduces c_validation/results.csv. Its unlocked phase contains one corruption record, six panic/trap records, two test failures, and one timeout.')
c_rows=result_rows(ROOT/'member1_completed'/'c_validation'/'results.csv','c_validation')
add_table('Table 16. C-runner dataset, attempt-level results', ['Run ID','Phase','Primary outcome','Elapsed (s)','Raw log'],c_rows,[.62,.82,1.35,.78,2.6],9.0)
p('The original set and C-runner set used the same xv6 revision, CPUs, RAM, guest test parameters, and phase counts. The C-runner helper changed how the host experiment was automated. The distinct unlocked outcome counts show why the two evidence sets must remain identifiable.')

heading('10. Conclusion, applications and future scope',1)
p('The case study connects race conditions to a shared kernel invariant. xv6’s kmem.lock serializes changes to the physical-page free list in kalloc() and kfree(). With the lock present, both recorded datasets passed all locked attempts. After the four allocator lock calls were temporarily removed, both sets recorded failures in all unlocked attempts. Restoring the allocator produced ten PASS outcomes in each restored phase.')
p('The results demonstrate the practical purpose of mutual exclusion under the specified QEMU workload. They do not prove correctness for every schedule or quantify the exact cause of each panic or corrupted byte. The test does not check every possible allocator invariant, and the removal also changes local interrupt-handling behavior associated with the spinlock calls.')
p('The main learning is that protection must cover the full read and update sequence that preserves an invariant. The same reasoning applies to kernel queues, process tables, buffer caches, counters, and shared application state. The inventory in this report is source-level analysis except for the physical allocator.')
p('Future work could vary the CPU count and stress parameters, add allocator-specific checks that detect lost or duplicate physical pages, examine another shared structure under a separate protocol, and measure lock contention. Those experiments were not performed here.')

heading('References',1)
references=[
    ('xv6 library (MIT PDOS)','https://github.com/mit-pdos/xv6-riscv'),
    ('MIT xv6 RISC-V book — Locking chapter','https://mit-pdos.github.io/xv6-riscv-book/lock.html'),
    ('The Open Group Base Specifications — memory synchronization','https://pubs.opengroup.org/onlinepubs/9699919799/basedefs/V1_chap04.html'),
    ('Michael and Scott — Nonblocking algorithms and the ABA problem','https://www.cs.rochester.edu/u/scott/papers/1998_JPDC_nonblocking.pdf')
]
for label,url in references:
    para=doc.add_paragraph(style='List Bullet')
    para.paragraph_format.space_after=Pt(5)
    para.paragraph_format.line_spacing=1.0
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    para.add_run(label+': ')
    add_hyperlink(para,url,url)

p('Evidence note: the terminal figures are static renders reconstructed from saved .cast recording output; they are not live desktop screenshots, and this report and deck contain no video. The recorded transcripts and results.csv remain in the supplied evidence directories.')
doc.core_properties.title='Case Study 14: Race Condition in XV6 Shared Data'
doc.core_properties.subject='Operating Systems case study report'
doc.core_properties.author='Operating Systems case study team'
doc.core_properties.keywords='xv6, race condition, synchronization, spinlock, allocator'
doc.save(DOCX_PATH)
print(DOCX_PATH)
print('original_attempts',len(orig_rows),'c_validation_attempts',len(c_rows))



