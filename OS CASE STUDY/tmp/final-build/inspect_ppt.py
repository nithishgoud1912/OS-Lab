from pathlib import Path
from pptx import Presentation
root=Path(r'C:\Users\HP\Documents\OS record\OS CASE STUDY')
ppt=root/'output/final/Case_Study_14_Race_Conditions_xv6_Final_v8.pptx'
prs=Presentation(str(ppt))
for i,s in enumerate(prs.slides,1):
    print(f'--- {i} ---')
    print(' | '.join(sh.text.replace('\\n',' / ') for sh in s.shapes if getattr(sh,'has_text_frame',False) and sh.text.strip())[:1500])
for name in ['A01-locked.cast','B02-unlocked.cast','C01-restored.cast']:
    p=root/'member1_completed/evidence'/name
    print('CAST',name,p.exists(),p.stat().st_size if p.exists() else None)
    if p.exists():
        lines=p.read_text(encoding='utf-8',errors='replace').splitlines()
        for line in lines[:4]: print(line[:1000])
