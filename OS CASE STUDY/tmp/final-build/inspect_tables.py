from pptx import Presentation
p=Presentation(r'C:\Users\HP\Documents\OS record\OS CASE STUDY\output\submission\Race_Conditions_and_Synchronization_xv6_Presentation.pptx')
for i in [2,4,17]:
 for sh in p.slides[i].shapes:
  if sh.has_table: print(i+1,'rows',len(sh.table.rows),'cols',len(sh.table.columns),[[c.text for c in row.cells] for row in sh.table.rows])
