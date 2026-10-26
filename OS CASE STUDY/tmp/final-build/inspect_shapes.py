from pathlib import Path
from pptx import Presentation
p=Presentation(r'C:\Users\HP\Documents\OS record\OS CASE STUDY\output\submission\Race_Conditions_and_Synchronization_xv6_Presentation.pptx')
for i in [4,5,6,7,16,17]:
 s=p.slides[i]; print('SLIDE',i+1)
 for sh in s.shapes:
  if getattr(sh,'has_text_frame',False) and sh.text.strip(): print('TXT',sh.text.replace('\\n',' / '))
  if sh.has_table: print('TABLE',len(sh.table.rows),len(sh.table.columns),[[c.text for c in row.cells] for row in sh.table.rows])
