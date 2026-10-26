from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_DATA_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pathlib import Path
import sys
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
pptx_path=Path(sys.argv[1])
prs=Presentation(str(pptx_path))
slide=prs.slides[17]
for shape in list(slide.shapes):
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        shape._element.getparent().remove(shape._element)
data=CategoryChartData()
data.categories=['Original set','C-runner']
for name,values,color in [
    ('Corruption',(3,1),'1F456E'),
    ('Panic / trap',(4,6),'405266'),
    ('Test failure',(2,2),'526477'),
    ('Timeout',(1,1),'65798B'),
]:
    data.add_series(name,values)
shape=slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(5.9), Inches(1.65), Inches(6.67), Inches(3.33), data)
chart=shape.chart
chart.has_title=True
chart.chart_title.text_frame.text='Unlocked outcomes, 10 attempts per dataset'
for paragraph in chart.chart_title.text_frame.paragraphs:
    for run in paragraph.runs:
        run.font.name='Arial'; run.font.size=Pt(16); run.font.bold=True; run.font.color.rgb=RGBColor(23,54,93)
chart.has_legend=True
chart.legend.position=XL_LEGEND_POSITION.BOTTOM
chart.legend.font.name='Arial'; chart.legend.font.size=Pt(14); chart.legend.font.color.rgb=RGBColor(83,97,110)
plot=chart.plots[0]
plot.overlap=100
plot.gap_width=55
plot.has_data_labels=True
labels=plot.data_labels
labels.position=XL_DATA_LABEL_POSITION.CENTER
labels.number_format='0'
labels.font.name='Arial'; labels.font.size=Pt(14); labels.font.bold=True; labels.font.color.rgb=RGBColor(255,255,255)
for series,color in zip(chart.series,['1F456E','405266','526477','65798B']):
    series.format.fill.solid(); series.format.fill.fore_color.rgb=RGBColor.from_string(color)
chart.value_axis.minimum_scale=0
chart.value_axis.maximum_scale=10
chart.value_axis.major_unit=2
chart.value_axis.has_major_gridlines=False
chart.value_axis.tick_labels.font.name='Arial'; chart.value_axis.tick_labels.font.size=Pt(14); chart.value_axis.tick_labels.font.color.rgb=RGBColor(83,97,110)
chart.category_axis.tick_labels.font.name='Arial'; chart.category_axis.tick_labels.font.size=Pt(14); chart.category_axis.tick_labels.font.color.rgb=RGBColor(23,32,42)
chart.category_axis.has_major_gridlines=False
cnv=shape._element.xpath('.//p:cNvPr')
if cnv: cnv[0].set('descr','Editable stacked bar chart comparing unlocked outcomes in the original and C-runner datasets.')
prs.save(str(pptx_path))
# Give chart-level default text an explicit family too; otherwise Office's
# inherited theme silently supplies Aptos and violates the deck's Arial policy.
tmp_path=pptx_path.with_suffix('.patched.pptx')
with ZipFile(pptx_path,'r') as zin, ZipFile(tmp_path,'w',ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        payload=zin.read(item.filename)
        if item.filename.startswith('ppt/charts/chart') and item.filename.endswith('.xml'):
            root=etree.fromstring(payload)
            ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
            for defrpr in root.xpath('.//c:txPr/a:p/a:pPr/a:defRPr | .//c:txPr/a:lstStyle/a:lvl1pPr/a:defRPr | .//c:txPr/a:lstStyle/a:defPPr/a:defRPr',namespaces=ns):
                latin=defrpr.find('{%s}latin'%ns['a'])
                if latin is None:
                    latin=etree.SubElement(defrpr,'{%s}latin'%ns['a'])
                latin.set('typeface','Arial')
            for axis_id in root.xpath('.//c:axId | .//c:crossAx',namespaces=ns):
                value=int(axis_id.get('val'))
                if value < 0:
                    axis_id.set('val',str(abs(value)))
            payload=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
        zout.writestr(item,payload)
tmp_path.replace(pptx_path)
print('Added editable stacked chart to slide 18 with embedded workbook.')
