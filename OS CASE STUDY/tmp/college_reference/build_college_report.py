from pathlib import Path
from copy import deepcopy
import re,json,hashlib
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import PageBreak,Paragraph,Preformatted,Table,Spacer

ROOT=Path(__file__).resolve().parents[2]
REF=Path(r'C:\Users\HP\Documents\OS record\Exp-1.docx')
TMP=ROOT/'tmp/college_reference'
OUT=ROOT/'output/college_format';OUT.mkdir(parents=True,exist_ok=True)
assert hashlib.sha256(REF.read_bytes()).hexdigest()=='c9f2993214ecb89d89a96fa6de5d26055cc74eac3dafb57b3d4628261d88be3c'
with ZipFile(REF) as z:
 (TMP/'package-inventory.json').write_text(json.dumps({n:{'size':len(z.read(n)),'sha256':hashlib.sha256(z.read(n)).hexdigest()} for n in z.namelist()},indent=2))

# Recover semantic flowables from the existing report authoring source without
# running its PDF writer or changing the previously delivered artifacts.
source=(ROOT/'tmp/build_submission.py').read_text(encoding='utf8')
ns={'__file__':str(ROOT/'tmp/build_submission.py')}
exec(compile(source[:source.index('def footer(')],'report_content','exec'),ns)
sections=[];part=[]
for obj in ns['story']:
 if isinstance(obj,PageBreak):sections.append(part);part=[]
 else:part.append(obj)
sections.append(part)
assert len(sections)==28

doc=Document(REF)
body=doc._element.body
record_section=deepcopy(doc.sections[-1]._sectPr)
for ch in list(body):body.remove(ch)
body.append(record_section)
sec=doc.sections[0]
sec.top_margin=Inches(.9);sec.bottom_margin=Inches(.8)
sec.left_margin=Inches(1);sec.right_margin=Inches(1)
sec.header_distance=Inches(.36);sec.footer_distance=Inches(.25)
sec.different_first_page_header_footer=False
doc.settings.odd_and_even_pages_header_footer=False
for r in list(sec._sectPr):
 if r.tag==qn('w:headerReference') and r.get(qn('w:type'))!='default':sec._sectPr.remove(r)
sec._sectPr.remove(sec._sectPr.find(qn('w:docGrid')))

def font(run,size=12,bold=None,name='Times New Roman'):
 run.font.name=name;run.font.size=Pt(size);run.font.color.rgb=RGBColor(0,0,0)
 if bold is not None:run.bold=bold
 rp=run._element.get_or_add_rPr();rf=rp.rFonts
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),name)
 for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  if qn('w:'+k) in rf.attrib:del rf.attrib[qn('w:'+k)]

for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(20 if name=='Title' else 14 if name=='Heading 1' else 12)
 pf=st.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(8);pf.line_spacing=Pt(16)
 if name.startswith('Heading'):
  pf.keep_with_next=True
  outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'0' if name=='Heading 1' else '1');st.element.get_or_add_pPr().append(outline)
 for e in st.element.xpath('.//w:color'):
  e.set(qn('w:val'),'000000')
  for k in ['themeColor','themeTint','themeShade']:e.attrib.pop(qn('w:'+k),None)
 for e in st.element.xpath('.//w:pBdr'):e.getparent().remove(e)
for name in ['TOC 1','TOC 2']:
 st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(2);st.paragraph_format.line_spacing=Pt(15)
front=doc.styles.add_style('Front Matter Heading',WD_STYLE_TYPE.PARAGRAPH)
front.base_style=doc.styles['Normal'];front.font.name='Times New Roman';front.font.size=Pt(14);front.font.bold=True

# Reuse the retained header, watermark, and footer rather than drawing a logo.
header=sec.header
for c,txt in zip(header.tables[0].rows[0].cells,['Operating Systems','Case Study 14','Date  __________']):
 c.text=txt
 for p in c.paragraphs:
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
  for r in p.runs:font(r,11)
footer=sec.footer
ft=footer.tables[0]
for c in ft.rows[0].cells:
 for p in list(c.paragraphs):p._element.getparent().remove(p._element)
p=ft.cell(0,0).add_paragraph();p.paragraph_format.space_after=Pt(0);font(p.add_run('Page No  '),11)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);fld.append(r);p._p.append(fld)
p=ft.cell(0,1).add_paragraph();p.paragraph_format.space_after=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
font(p.add_run('Signature of the Faculty  ______________'),11)
for p in list(footer.paragraphs):
 p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(1)
 if not p.text:
  for r in p.runs:
   if not r._r.xpath('.//w:drawing') and not r._r.xpath('.//w:pict'):font(r,1)

def para(txt,style=None,align=None,size=12,bold=None):
 p=doc.add_paragraph(style=style)
 p.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
 p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=Pt(16)
 font(p.add_run(txt),size,bold)
 return p
def heading(txt,level=1):
 txt=re.sub(r'^(\d+)\.\s+',r'\1 ',txt).replace(':','').replace('(','').replace(')','')
 if txt in ['Certificate and approval details','Members and contributions','Index']:
  return para(txt,'Front Matter Heading',WD_ALIGN_PARAGRAPH.LEFT,14,True)
 return para(txt,'Heading '+str(level),WD_ALIGN_PARAGRAPH.LEFT,14 if level==1 else 12,True)
def code(txt):
 p=para(txt,align=WD_ALIGN_PARAGRAPH.LEFT,size=9.5)
 p.paragraph_format.line_spacing=Pt(12.5);p.paragraph_format.space_after=Pt(9)
 p.paragraph_format.keep_together=True
 font(p.runs[0],9.5,name='Courier New')
 return p
def table(data,widths=None):
 t=doc.add_table(rows=len(data),cols=len(data[0]));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 width=sec.page_width-sec.left_margin-sec.right_margin
 if widths is None:widths=[1]*len(data[0])
 widths=[width*x/sum(widths) for x in widths]
 for i,w in enumerate(widths):t.columns[i].width=int(w)
 pr=t._tbl.tblPr
 bd=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
 for e in pr.findall(qn('w:tblBorders')):pr.remove(e)
 pr.append(bd)
 margins=OxmlElement('w:tblCellMar')
 for k,v in [('top',65),('bottom',65),('left',85),('right',85)]:
  e=OxmlElement('w:'+k);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');margins.append(e)
 pr.append(margins)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ri==0:trpr.append(OxmlElement('w:tblHeader'))
  for ci,(cell,txt) in enumerate(zip(row.cells,data[ri])):
   cell.width=int(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cell.text=str(txt);p=cell.paragraphs[0]
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci>0 and len(str(txt))<16 else WD_ALIGN_PARAGRAPH.LEFT
   p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(13)
   for r in p.runs:font(r,10.5,ri==0)
   shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EDF2' if ri==0 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(shading)
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
 return t
def field(instruction):
 p=doc.add_paragraph();r=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
 instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=instruction;r._r.append(instr)
 sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
 font(p.add_run(''),11)
 end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');p.add_run()._r.append(end)

def to_plain(x):
 if isinstance(x,Paragraph):return x.getPlainText().replace('<br/>','\n')
 return str(x)
def render_section(items):
 first=True
 for obj in items:
  if isinstance(obj,Preformatted):code('\n'.join(obj.lines))
  elif isinstance(obj,Paragraph):
   txt=obj.getPlainText()
   if first:heading(txt);first=False;continue
   if obj.style.name=='head':heading(txt)
   elif obj.style.name=='sub':para(txt,align=WD_ALIGN_PARAGRAPH.LEFT,bold=True).paragraph_format.keep_with_next=True
   elif obj.style.name=='caption':
    p=para(txt,'Caption',WD_ALIGN_PARAGRAPH.LEFT,10);p.paragraph_format.line_spacing=Pt(13);p.runs[0].italic=True
   else:para(txt)
  elif isinstance(obj,Table):
   dat=[[to_plain(c) for c in row] for row in obj._cellvalues]
   table(dat,obj._colWidths)
  elif isinstance(obj,Spacer):
   # Flowable padding is represented by paragraph/table spacing in Word.
   pass

contributions=[
 ['L Nithish Kumar Goud\n160124733327','Hands-on experiment and results','Build and run xv6; prepare the C stress test; compare original, removed and restored locks; record results and restoration evidence.','25%'],
 ['G Ishan\n160124733322','Race analysis and interleavings','Analyse kalloc and kfree races; explain possible incorrect outcomes; compare a scheduler or PID scenario and relate theory to results.','25%'],
 ['Santhosh\n160124733332','Locking mechanisms and critical sections','Explain spinlock internals, atomic operations, memory ordering and interrupt handling; identify protected sections and lock discipline.','25%'],
 ['Venkat\n160124733340','Foundations and report integration','Explain core terms and shared-data inventory; compile and format the report; review evidence, limitations and references.','25%']
]
for idx,items in enumerate(sections):
 if idx:doc.add_page_break()
 if idx==0:
  p=para('OPERATING SYSTEMS CASE STUDY','Title',WD_ALIGN_PARAGRAPH.CENTER,20,True);p.paragraph_format.space_before=Pt(45);p.paragraph_format.space_after=Pt(20)
  para('Race Conditions and Synchronization in xv6','Title',WD_ALIGN_PARAGRAPH.CENTER,20,True)
  para('Physical page allocator experiment and synchronization analysis',align=WD_ALIGN_PARAGRAPH.CENTER)
  para('Chaitanya Bharathi Institute of Technology',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True)
  para('Department of Computer Science and Engineering',align=WD_ALIGN_PARAGRAPH.CENTER)
  para('Submitted by',align=WD_ALIGN_PARAGRAPH.LEFT,bold=True).paragraph_format.space_before=Pt(28)
  table([['Name','Roll number']]+ns['members'],[3,2])
  for x in ['Faculty  ______________________________','Semester and Section  __________________','Academic Year  ________________________']:
   para(x,align=WD_ALIGN_PARAGRAPH.LEFT)
 elif idx==1:
  heading('Certificate and approval details')
  para('Race Conditions and Synchronization in xv6',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True)
  para('Chaitanya Bharathi Institute of Technology\nDepartment of Computer Science and Engineering',align=WD_ALIGN_PARAGRAPH.CENTER)
  table([['Field','Details'],['Course','Operating Systems'],['Faculty and designation','____________________________'],['Semester and section','____________________________'],['Academic year','____________________________']],[1.5,3.5])
  para('Certificate',bold=True)
  for _ in range(6):para('____________________________________________________________',align=WD_ALIGN_PARAGRAPH.LEFT)
  table([['Faculty signature','Department approval','Date'],['________________','________________','____________']],[2,2,1])
 elif idx==2:
  heading('Members and contributions')
  para('The case study is divided equally among four members. Each member has a 25 percent responsibility share and leads the section described below. The share represents the agreed division of work rather than measured hours.')
  para('Table 1 Team contributions',align=WD_ALIGN_PARAGRAPH.LEFT,bold=True)
  table([['Member and roll number','Primary contribution','Responsibility','Share']]+contributions,[1.7,1.5,3.3,.55])
  para('Team integration',bold=True)
  para('Nithish supplies the practical evidence and results for Ishan to compare with the possible race scenarios. Santhosh verifies the protecting locks and critical sections. Venkat integrates the explanations, checks source references and result consistency, and prepares the final report. All four members review the complete experiment and its limitations.')
 elif idx==3:
  heading('Index')
  field('TOC \\o "1-1" \\h \\z \\u')
 else:render_section(items)
 if idx==7:
  doc.add_page_break();heading('Shared kernel data and protecting locks')
  para('This source-level inventory connects each shared resource to its synchronization rule. It supports the foundations and lock-analysis responsibilities. Only the physical-page allocator was modified in the measured experiment; the other structures are included for comparison.')
  para('Table A Shared data inventory',align=WD_ALIGN_PARAGRAPH.LEFT,bold=True)
  table([['Shared resource','Source file','Protection','Reason'],['Physical free list','kernel/kalloc.c','kmem.lock','Preserve page-list push and pop updates.'],['Process state and context','kernel/proc.c','p->lock','Coordinate state transitions and scheduling.'],['Next process ID','kernel/proc.c','pid_lock','Serialize allocation of a unique PID.'],['Parent child relationships','kernel/proc.c','wait_lock','Coordinate parent relationships and wait/wakeup.'],['Buffer cache metadata','kernel/bio.c','bcache.lock','Protect lookup, references and cache-list updates.'],['Open file table','kernel/file.c','ftable.lock','Protect file-object reference accounting.'],['Inode table metadata','kernel/fs.c','itable.lock','Protect inode references and table lookup; inode contents use sleep-locks.'],['File-system log state','kernel/log.c','log.lock','Coordinate transaction bookkeeping.'],['Tick counter','kernel/trap.c','tickslock','Serialize shared tick updates and related waits.'],['Pipe state and buffer','kernel/pipe.c','pi->lock','Coordinate readers, writers and buffer counters.']],[1.4,1.35,1.05,2.5])
  para('A shared structure may need more than one kind of protection. For example, the cache metadata lock differs from the sleep-lock protecting an individual buffer. The inventory describes the stated responsibility of each listed lock rather than claiming it protects every field in its source file.')
 if idx==15:
  doc.add_page_break();heading('Scheduler and PID race analysis')
  para('The scheduler provides a second shared-state example for race analysis. In the examined source, scheduler() acquires p->lock before checking whether p->state is RUNNABLE. It changes the state to RUNNING and switches context under the process-lock protocol. The process and scheduler coordinate release and reacquisition across the context switch.')
  para('Table B Possible scheduler selection without p lock',align=WD_ALIGN_PARAGRAPH.LEFT,bold=True)
  table([['Step','CPU 0','CPU 1'],['1','Observes process P as RUNNABLE',''],['2','','Observes the same P as RUNNABLE'],['3','Marks P RUNNING and selects it',''],['4','','Also marks P RUNNING and selects it']],[.5,2.8,2.8])
  para('Without the required coordination, this hypothetical schedule can let two CPUs select one process and interfere with the same process context or kernel stack. With the protocol, checking and changing the state cannot be performed as two independent unprotected selections. The lock handoff across swtch() is part of the mechanism and should not be reduced to an ordinary user-space lock example.')
  para('PID allocation',bold=True)
  code('acquire(&pid_lock);\npid = nextpid;\nnextpid = nextpid + 1;\nrelease(&pid_lock);')
  para('The PID counter is another read-modify-write sequence. If two CPUs read the same nextpid without the lock, both could select the same identifier. pid_lock serializes that sequence. These scheduler and PID scenarios are source analysis only; neither lock was removed or tested in the allocator experiment.')

# All section references are headed with Word's real styles so the TOC can be
# refreshed by Word. Preserve the image and recurring VML component exactly.
settings=doc.settings.element
up=settings.find(qn('w:updateFields'))
if up is None:up=OxmlElement('w:updateFields');settings.append(up)
up.set(qn('w:val'),'true')
doc.core_properties.title='Race Conditions and Synchronization in xv6'
doc.core_properties.subject='Operating Systems case study'
doc.core_properties.author='L Nithish Kumar Goud; G Ishan; Santhosh; Venkat'
final=OUT/'Race_Conditions_and_Synchronization_xv6_College_Report.docx'
doc.save(final)
with ZipFile(REF) as a,ZipFile(final) as b:
 assert a.read('word/media/image1.png')==b.read('word/media/image1.png')
 unchanged=[n for n in a.namelist() if n in b.namelist() and a.read(n)==b.read(n)]
 (TMP/'preservation-check.json').write_text(json.dumps({'unchanged_parts':unchanged,'unchanged_logo':True,'intentional_changes':['document body and section layout','styles for required report typography','header record metadata','footer page field','field refresh settings','core report metadata']},indent=2))
print(final)
